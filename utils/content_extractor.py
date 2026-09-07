"""
Content extraction for AI-suggested filenames.

Given a file, produces either:
  - a short text preview (for text-ish documents), or
  - raw image bytes + media type (for pictures),
so utils/ai_namer.py has something meaningful to show the model.

Every extractor here is best-effort and silent on failure — if a file can't
be read (corrupt, password-protected, unsupported internal format, etc.)
this just returns None and the caller falls back to the normal rename
convention. Nothing here ever raises.
"""

from pathlib import Path
from typing import Optional, Tuple

from config import settings

TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".log", ".json"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}

_IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
}


def _read_plain_text(file_path: Path, max_chars: int) -> Optional[str]:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(max_chars)
    except OSError:
        return None


def _read_pdf_text(file_path: Path, max_chars: int) -> Optional[str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    try:
        reader = PdfReader(str(file_path))
        chunks = []
        total = 0
        for page in reader.pages[:5]:  # first few pages is plenty for naming
            text = page.extract_text() or ""
            chunks.append(text)
            total += len(text)
            if total >= max_chars:
                break
        combined = "\n".join(chunks)[:max_chars]
        return combined if combined.strip() else None
    except Exception:
        return None  # encrypted / malformed PDFs, etc. — just skip AI naming


def _read_docx_text(file_path: Path, max_chars: int) -> Optional[str]:
    try:
        import docx
    except ImportError:
        return None
    try:
        doc = docx.Document(str(file_path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return text[:max_chars] or None
    except Exception:
        return None


def get_text_preview(file_path: Path) -> Optional[str]:
    """Returns a short text preview of the file's content, or None if the
    file type isn't a supported text-ish format or reading failed."""
    max_chars = settings.AI_RENAME_PREVIEW_CHARS
    ext = file_path.suffix.lower()

    if ext in TEXT_EXTENSIONS:
        return _read_plain_text(file_path, max_chars)
    if ext in PDF_EXTENSIONS:
        return _read_pdf_text(file_path, max_chars)
    if ext in DOCX_EXTENSIONS:
        return _read_docx_text(file_path, max_chars)
    return None


def _rasterize_pdf_first_page(
    file_path: Path, max_bytes: int
) -> Optional[Tuple[str, bytes]]:
    """Renders page 1 of a PDF to PNG bytes, for scanned/image-only PDFs
    where _read_pdf_text() found no usable text layer. Lets Gemini's vision
    input read the page the same way it reads a photo. Best-effort/silent:
    returns None on any failure (missing dependency, corrupt/encrypted PDF,
    zero-page PDF, oversized render, etc.)."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return None

    try:
        pdf = pdfium.PdfDocument(str(file_path))
        try:
            if len(pdf) == 0:
                return None
            page = pdf[0]
            # ~150 DPI (scale = dpi / 72) is plenty for a naming model to
            # read titles/headings without producing a huge payload.
            bitmap = page.render(scale=150 / 72)
            pil_image = bitmap.to_pil()
        finally:
            pdf.close()

        import io

        buf = io.BytesIO()
        pil_image.save(buf, format="PNG")
        data = buf.getvalue()
        if len(data) > max_bytes:
            return None
        return "image/png", data
    except Exception:
        return None  # encrypted / malformed / password-protected PDFs, etc.


def get_image_payload(file_path: Path) -> Optional[Tuple[str, bytes]]:
    """Returns (media_type, raw_bytes) for supported image types, capped at
    AI_RENAME_MAX_IMAGE_MB to avoid sending huge files to the API. None if
    unsupported or unreadable.

    Also handles PDFs as a fallback: if a PDF has no usable text layer
    (scanned/image-only pages), the first page is rasterized to PNG and
    sent to Gemini's vision input instead. This is only reached for PDFs
    when get_text_preview() already returned None."""
    ext = file_path.suffix.lower()
    max_bytes = settings.AI_RENAME_MAX_IMAGE_MB * 1024 * 1024

    if ext in PDF_EXTENSIONS:
        return _rasterize_pdf_first_page(file_path, max_bytes)

    if ext not in IMAGE_EXTENSIONS:
        return None

    try:
        if file_path.stat().st_size > max_bytes:
            return None
        data = file_path.read_bytes()
    except OSError:
        return None

    return _IMAGE_MEDIA_TYPES[ext], data


def is_ai_nameable(file_path: Path) -> bool:
    """Quick check used by the pipeline before doing any real work."""
    ext = file_path.suffix.lower()
    if ext not in settings.AI_RENAME_EXTENSIONS:
        return False
    return ext in (
        TEXT_EXTENSIONS | PDF_EXTENSIONS | DOCX_EXTENSIONS | IMAGE_EXTENSIONS
    )
